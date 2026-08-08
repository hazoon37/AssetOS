from __future__ import annotations

from io import BytesIO

import pandas as pd

from services.excel_asset_service import (
    EXCEL_COLUMNS,
    GUIDE_SHEET_NAME,
    REQUIRED_COLUMNS,
    TEMPLATE_COLUMNS,
    build_asset_template_excel,
    build_validation_error_report,
    export_assets_excel,
    validate_asset_excel,
)


def _excel_bytes(dataframe: pd.DataFrame) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name="Assets", index=False)
    return output.getvalue()


def test_validate_asset_excel_success() -> None:
    dataframe = pd.DataFrame([
        {
            "자산종류": "미국주식",
            "자산명": "Apple",
            "티커": "AAPL",
            "수량": 2,
            "평균단가": 180,
            "현재가": 210,
            "통화": "USD",
        }
    ])
    result = validate_asset_excel(_excel_bytes(dataframe))
    assert result.success is True
    assert len(result.rows) == 1
    assert result.rows[0]["symbol"] == "AAPL"


def test_validate_asset_excel_missing_required_column() -> None:
    dataframe = pd.DataFrame([{"자산명": "Apple"}])
    result = validate_asset_excel(_excel_bytes(dataframe))
    assert result.success is False
    assert result.errors


def test_validate_minimal_smart_import_row() -> None:
    dataframe = pd.DataFrame([{"자산명": "Apple", "수량": 2}])
    result = validate_asset_excel(_excel_bytes(dataframe))
    assert result.success is True
    assert result.rows[0]["asset_name"] == "Apple"
    assert result.rows[0]["quantity"] == 2.0
    assert result.rows[0]["_average_price_missing"] is True


def test_validate_minimal_row_requires_quantity_value() -> None:
    dataframe = pd.DataFrame([{"자산명": "Apple", "수량": None}])
    result = validate_asset_excel(_excel_bytes(dataframe))
    assert result.success is False
    assert any("수량이 비어" in error for error in result.errors)
    report = pd.read_excel(BytesIO(build_validation_error_report(result)))
    assert report.loc[0, "Status"] == "🔴 Fail"
    assert "2행" in report.loc[0, "Suggested Action"]


def test_validate_generic_excel_with_automatic_mapping() -> None:
    output = BytesIO()
    dataframe = pd.DataFrame([{
        "종목명": "Apple",
        "보유수량": 2,
        "보유금액": 360,
        "통화코드": "USD",
    }])
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name="보유자산", index=False)

    result = validate_asset_excel(output.getvalue())
    assert result.success is True
    assert result.rows[0]["asset_name"] == "Apple"
    assert result.rows[0]["quantity"] == 2.0
    assert result.rows[0]["average_price"] == 180.0
    assert result.mapped_columns["종목명"] == "자산명"


def test_template_matches_excel_schema() -> None:
    template_bytes = build_asset_template_excel()
    template = pd.read_excel(BytesIO(template_bytes), sheet_name="Assets")
    assert template.columns.tolist() == TEMPLATE_COLUMNS
    assert template.iloc[0]["자산명"] == "Apple"
    assert template.iloc[0]["Ticker"] == "AAPL"

    result = validate_asset_excel(template_bytes)
    assert result.success is True
    assert result.rows[0]["symbol"] == "AAPL"


def test_validation_preserves_original_excel_row_number() -> None:
    dataframe = pd.DataFrame([
        {"자산명": None, "수량": None},
        {"자산명": "Apple", "수량": 2},
    ])
    result = validate_asset_excel(_excel_bytes(dataframe))
    assert result.success is True
    assert result.rows[0]["_excel_row_number"] == 3


def test_template_workbook_ux() -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(build_asset_template_excel()))
    assert workbook.sheetnames == ["Assets", GUIDE_SHEET_NAME]

    assets = workbook["Assets"]
    assert assets.freeze_panes == "A2"
    assert assets.auto_filter.ref == assets.dimensions
    assert assets["A2"].font.italic is True
    assert assets["A2"].comment.text == "예시입니다. 삭제 후 사용하세요."
    assert len(assets.data_validations.dataValidation) >= 2

    headers = {cell.value: cell for cell in assets[1]}
    for required_column in REQUIRED_COLUMNS:
        assert headers[required_column].font.bold is True
        assert headers[required_column].fill.fgColor.rgb == "00D9EAF7"
        assert headers[required_column].comment.text == "* 필수 입력 항목"


def test_export_can_be_imported_without_losing_asset_fields() -> None:
    assets = pd.DataFrame([
        {
            "id": 1,
            "asset_type": "미국주식",
            "asset_name": "Apple",
            "symbol": "AAPL",
            "quantity": 2.0,
            "average_price": 180.0,
            "current_price": 210.0,
            "currency": "USD",
            "memo": "장기",
            "asset_class": "EQUITY",
            "country": "US",
            "exchange": "NASDAQ",
            "sector": "Technology",
            "industry": "Consumer Electronics",
            "is_cash": 0,
            "is_leverage": 0,
            "is_inverse": 0,
            "leverage_multiple": 1.0,
            "data_source": "Test",
            "tags": "core",
            "created_at": "2026-01-01 00:00:00",
            "updated_at": "2026-01-02 00:00:00",
        }
    ])
    exported = export_assets_excel(assets)
    workbook = pd.read_excel(BytesIO(exported), sheet_name="Assets")
    assert workbook.columns.tolist() == EXCEL_COLUMNS

    result = validate_asset_excel(exported)
    assert result.success is True
    assert result.rows[0]["asset_name"] == "Apple"
    assert result.rows[0]["exchange"] == "NASDAQ"


def test_duplicate_ticker_is_reported_without_merging_rows() -> None:
    dataframe = pd.DataFrame([
        {"자산명": "Apple lot 1", "티커": "AAPL", "수량": 1},
        {"자산명": "Apple lot 2", "티커": "AAPL", "수량": 2},
    ])
    result = validate_asset_excel(_excel_bytes(dataframe))
    assert result.success is True
    assert len(result.rows) == 2
    assert any("자동 합산하지 않습니다" in warning for warning in result.warnings)


def test_import_file_size_limit_has_friendly_error() -> None:
    result = validate_asset_excel(b"x" * (10 * 1024 * 1024 + 1))
    assert result.success is False
    assert any("10MB" in error for error in result.errors)
