from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import re
from typing import Any, BinaryIO

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from components.asset_management.constants import ASSET_TYPES, SUPPORTED_CURRENCIES
from repositories import get_asset_repository
from repositories.asset_repository import AssetRepository
from services.import_detector import ImportFileType, detect_import_file
from services.import_mapping import map_import_columns

SHEET_NAME = "Assets"
REQUIRED_COLUMNS = {
    "자산명",
    "수량",
}
TEMPLATE_COLUMNS = ["자산종류", "자산명", "Ticker", "수량", "평균단가", "현재가", "통화"]

COLUMN_MAP = {
    "ID": "id",
    "자산종류": "asset_type",
    "자산명": "asset_name",
    "티커": "symbol",
    "수량": "quantity",
    "평균단가": "average_price",
    "현재가": "current_price",
    "통화": "currency",
    "메모": "memo",
    "자산군": "asset_class",
    "국가": "country",
    "거래소": "exchange",
    "섹터": "sector",
    "산업": "industry",
    "현금성": "is_cash",
    "레버리지": "is_leverage",
    "인버스": "is_inverse",
    "레버리지배수": "leverage_multiple",
    "데이터출처": "data_source",
    "태그": "tags",
    "등록일": "created_at",
    "수정일": "updated_at",
}

EXCEL_COLUMNS = list(COLUMN_MAP)
GUIDE_SHEET_NAME = "📖 작성가이드"

TRUE_VALUES = {"true", "1", "yes", "y", "예", "네", "o"}
FALSE_VALUES = {"false", "0", "no", "n", "아니오", "아니요", "x", ""}
MAX_IMPORT_FILE_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000


@dataclass(frozen=True)
class ExcelValidationResult:
    success: bool
    rows: list[dict[str, Any]]
    preview: pd.DataFrame
    errors: list[str]
    warnings: list[str]
    file_type: ImportFileType | None = None
    mapped_columns: dict[str, str] = field(default_factory=dict)
    source_dataframe: pd.DataFrame = field(
        default_factory=pd.DataFrame, repr=False, compare=False
    )


def _clean_text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def _to_float(value: Any, *, field: str, row_number: int) -> float:
    if value is None or pd.isna(value) or str(value).strip() == "":
        return 0.0
    try:
        return float(str(value).replace(",", "").replace("₩", "").strip())
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{row_number}행의 '{field}' 값이 숫자가 아닙니다: {value}") from exc


def _to_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None or pd.isna(value):
        return False
    normalized = str(value).strip().lower()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    return False


def read_asset_excel(file: BinaryIO | bytes, sheet_name: str = SHEET_NAME) -> pd.DataFrame:
    source = BytesIO(file) if isinstance(file, bytes) else file
    return pd.read_excel(source, sheet_name=sheet_name, engine="openpyxl")


def _dataframe_to_excel(dataframe: pd.DataFrame) -> bytes:
    """Serialize one Assets sheet without touching the filesystem."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name=SHEET_NAME, index=False)
    return output.getvalue()


def build_asset_template_excel() -> bytes:
    """Return a styled, importer-compatible workbook for entering assets."""
    example = {
        "자산종류": "미국주식",
        "자산명": "Apple",
        "Ticker": "AAPL",
        "수량": 2,
        "평균단가": 180,
        "현재가": 210,
        "통화": "USD",
    }
    template = pd.DataFrame([example], columns=TEMPLATE_COLUMNS)
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        template.to_excel(writer, sheet_name=SHEET_NAME, index=False)
        workbook = writer.book
        assets_sheet = writer.sheets[SHEET_NAME]

        required_fill = PatternFill("solid", fgColor="D9EAF7")
        optional_fill = PatternFill("solid", fgColor="EAF0F6")
        example_fill = PatternFill("solid", fgColor="E7E7E7")

        for cell in assets_sheet[1]:
            cell.font = Font(bold=True, color="1F2937")
            cell.fill = required_fill if cell.value in REQUIRED_COLUMNS else optional_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.value in REQUIRED_COLUMNS:
                cell.comment = Comment("* 필수 입력 항목", "AssetOS")

        for cell in assets_sheet[2]:
            cell.fill = example_fill
            cell.font = Font(italic=True, color="666666")
        assets_sheet["A2"].comment = Comment(
            "예시입니다. 삭제 후 사용하세요.",
            "AssetOS",
        )

        assets_sheet.freeze_panes = "A2"
        assets_sheet.auto_filter.ref = assets_sheet.dimensions
        assets_sheet.row_dimensions[1].height = 24

        dropdowns = {
            "자산종류": ASSET_TYPES,
            "통화": SUPPORTED_CURRENCIES,
            "현금성": ["예", "아니오"],
            "레버리지": ["예", "아니오"],
            "인버스": ["예", "아니오"],
        }
        for column_name, choices in dropdowns.items():
            if column_name not in TEMPLATE_COLUMNS:
                continue
            column_index = TEMPLATE_COLUMNS.index(column_name) + 1
            formula = '"' + ",".join(choices) + '"'
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            validation.error = "목록에서 값을 선택해 주세요."
            validation.errorTitle = "올바르지 않은 값"
            validation.prompt = "목록에서 값을 선택할 수 있습니다."
            validation.promptTitle = column_name
            validation.showErrorMessage = True
            validation.showInputMessage = True
            assets_sheet.add_data_validation(validation)
            validation.add(f"{get_column_letter(column_index)}2:{get_column_letter(column_index)}1000")

        for column_index, column_name in enumerate(TEMPLATE_COLUMNS, start=1):
            values = [column_name, example.get(column_name, "")]
            width = min(max(len(str(value)) for value in values) + 4, 32)
            assets_sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 11)

        guide = workbook.create_sheet(GUIDE_SHEET_NAME)
        guide_rows = [
            ("AssetOS Excel 작성가이드", ""),
            ("사용 방법", "Assets 시트의 회색 예시 행을 삭제한 뒤 자산을 한 행씩 입력하세요."),
            ("필수 항목 (*)", "자산명 *, 수량 *"),
            ("티커 (선택)", "비워 두면 자산명으로 자동 조회합니다. 입력하면 티커를 먼저 검증하고 사용합니다."),
            ("자동 조회", "자산종류, 티커, 현재가, 통화, 거래소, 국가는 업로드 후 가능한 경우 자동으로 조회합니다."),
            ("선택 입력", "평균단가는 선택 사항입니다. 입력하면 평가손익과 수익률을 계산할 수 있습니다."),
            ("업로드 과정", "Asset Manager에서 파일 선택 → 미리보기 및 오류 확인 → 동의 → Apply Changes"),
            ("주의", "열 이름과 Assets 시트 이름을 변경하지 마세요. 적용 전 현재 DB가 자동 백업됩니다."),
            ("간단한 예", "미국주식 | Apple | AAPL | 2 | 180 | 210 | USD"),
        ]
        for row in guide_rows:
            guide.append(row)
        guide["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        guide["B1"].font = Font(size=16, bold=True, color="FFFFFF")
        for cell in guide[1]:
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for row_index in range(2, guide.max_row + 1):
            guide.cell(row_index, 1).font = Font(bold=True, color="1F2937")
            guide.cell(row_index, 1).fill = required_fill
            guide.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
        guide.column_dimensions["A"].width = 20
        guide.column_dimensions["B"].width = 86
        guide.freeze_panes = "A2"

    return output.getvalue()


def export_assets_excel(
    assets: pd.DataFrame | None = None,
    repository: AssetRepository | None = None,
) -> bytes:
    """Export the current assets table using the standard import column names."""
    source = (repository or get_asset_repository()).get_assets() if assets is None else assets
    exported = pd.DataFrame()
    for excel_column, database_column in COLUMN_MAP.items():
        exported[excel_column] = (
            source[database_column]
            if database_column in source.columns
            else pd.Series(index=source.index, dtype="object")
        )
    return _dataframe_to_excel(exported)


def validate_asset_excel(file: BinaryIO | bytes) -> ExcelValidationResult:
    errors: list[str] = []
    warnings: list[str] = []

    try:
        if isinstance(file, bytes):
            file_bytes = file
        else:
            file.seek(0)
            file_bytes = file.read()
            file.seek(0)
        if len(file_bytes) > MAX_IMPORT_FILE_BYTES:
            raise ValueError("Excel 파일은 10MB 이하여야 합니다.")
        detection = detect_import_file(file_bytes)
        dataframe = read_asset_excel(file_bytes, detection.sheet_name)
        source_dataframe = dataframe.copy()
        if len(dataframe) > MAX_IMPORT_ROWS:
            raise ValueError(f"한 번에 최대 {MAX_IMPORT_ROWS:,}행까지 가져올 수 있습니다.")
        mapping = map_import_columns(dataframe)
        dataframe = mapping.dataframe
        warnings.extend(mapping.warnings)
    except Exception as exc:
        return ExcelValidationResult(
            success=False,
            rows=[],
            preview=pd.DataFrame(),
            errors=[f"엑셀 파일을 읽지 못했습니다: {exc}"],
            warnings=[],
        )

    dataframe.columns = [str(column).strip() for column in dataframe.columns]
    missing = sorted(REQUIRED_COLUMNS - set(dataframe.columns))
    if missing:
        missing_preview = dataframe.head(20).copy()
        missing_preview.insert(
            0, "Row Number", [int(index) + 2 for index in missing_preview.index]
        )
        return ExcelValidationResult(
            success=False,
            rows=[],
            preview=missing_preview,
            errors=["필수 열이 없습니다: " + ", ".join(missing)],
            warnings=warnings,
            file_type=detection.file_type,
            mapped_columns=mapping.mappings,
            source_dataframe=source_dataframe,
        )

    # Completely empty rows are ignored.
    dataframe = dataframe.dropna(how="all").copy()
    rows: list[dict[str, Any]] = []

    for index, excel_row in dataframe.iterrows():
        row_number = int(index) + 2
        asset_name = _clean_text(excel_row.get("자산명"))
        asset_type = _clean_text(excel_row.get("자산종류"))

        # Template spare rows are ignored.
        if not asset_name and not _clean_text(excel_row.get("수량")):
            continue

        if not asset_name:
            errors.append(f"{row_number}행: 자산명이 비어 있습니다.")
            continue
        raw_quantity = excel_row.get("수량")
        if raw_quantity is None or pd.isna(raw_quantity) or str(raw_quantity).strip() == "":
            errors.append(f"{row_number}행 '{asset_name}': 수량이 비어 있습니다.")
            continue
        try:
            quantity = _to_float(raw_quantity, field="수량", row_number=row_number)
            average_price = _to_float(excel_row.get("평균단가"), field="평균단가", row_number=row_number)
            current_price = _to_float(excel_row.get("현재가"), field="현재가", row_number=row_number)
            leverage_multiple = _to_float(
                excel_row.get("레버리지배수", 1),
                field="레버리지배수",
                row_number=row_number,
            ) or 1.0
        except ValueError as exc:
            errors.append(str(exc))
            continue

        if quantity < 0 or average_price < 0 or current_price < 0:
            errors.append(f"{row_number}행: 수량과 가격은 음수가 될 수 없습니다.")
            continue
        if quantity == 0:
            warnings.append(f"{row_number}행 '{asset_name}': 수량이 0입니다.")

        currency = _clean_text(excel_row.get("통화")).upper()

        row = {
            "_excel_row_number": row_number,
            "asset_type": asset_type,
            "asset_name": asset_name,
            "symbol": _clean_text(excel_row.get("티커")).upper(),
            "quantity": quantity,
            "average_price": average_price,
            "current_price": current_price,
            "currency": currency,
            "memo": _clean_text(excel_row.get("메모")),
            "asset_class": _clean_text(excel_row.get("자산군")),
            "country": _clean_text(excel_row.get("국가")).upper(),
            "exchange": _clean_text(excel_row.get("거래소")),
            "sector": _clean_text(excel_row.get("섹터")),
            "industry": _clean_text(excel_row.get("산업")),
            "is_cash": _to_bool(excel_row.get("현금성")),
            "is_leverage": _to_bool(excel_row.get("레버리지")),
            "is_inverse": _to_bool(excel_row.get("인버스")),
            "leverage_multiple": leverage_multiple,
            "data_source": _clean_text(excel_row.get("데이터출처")) or "Excel Import",
            "tags": _clean_text(excel_row.get("태그")),
            "_average_price_missing": (
                excel_row.get("평균단가") is None
                or pd.isna(excel_row.get("평균단가"))
                or str(excel_row.get("평균단가")).strip() == ""
            ),
        }
        rows.append(row)

    duplicate_keys: dict[tuple[str, str], list[str]] = {}
    for row in rows:
        symbol = str(row.get("symbol") or "").strip().upper()
        exchange = str(row.get("exchange") or "").strip().upper()
        if symbol:
            duplicate_keys.setdefault((symbol, exchange), []).append(str(row["asset_name"]))
    for (symbol, exchange), names in duplicate_keys.items():
        if len(names) > 1:
            market = f"/{exchange}" if exchange else ""
            warnings.append(
                f"동일 종목 {symbol}{market}이 {len(names):,}개 행에 있습니다. 자동 합산하지 않습니다."
            )

    if not rows and not errors:
        errors.append("가져올 자산 행이 없습니다.")

    preview_columns = [
        column
        for column in ["자산종류", "자산명", "티커", "수량", "평균단가", "현재가", "통화"]
        if column in dataframe.columns
    ]
    preview = dataframe[preview_columns].head(30).copy()
    preview.insert(0, "Row Number", [int(index) + 2 for index in preview.index])

    return ExcelValidationResult(
        success=not errors,
        rows=rows,
        preview=preview,
        errors=errors,
        warnings=warnings,
        file_type=detection.file_type,
        mapped_columns=mapping.mappings,
        source_dataframe=source_dataframe,
    )


def backup_database(repository: AssetRepository | None = None) -> Path | None:
    return (repository or get_asset_repository()).backup()


def build_validation_error_report(result: ExcelValidationResult) -> bytes:
    """Export validation failures while preserving the uploaded columns."""
    report = result.source_dataframe.copy()
    messages_by_row: dict[int, list[str]] = {}
    global_messages: list[str] = []
    for message in result.errors:
        match = re.match(r"^(\d+)행", message)
        if match:
            messages_by_row.setdefault(int(match.group(1)), []).append(message)
        else:
            global_messages.append(message)
    report["Status"] = "🔴 Fail"
    messages = [
        " · ".join(messages_by_row.get(int(index) + 2, global_messages))
        for index in report.index
    ]
    report["Suggested Action"] = [
        f"Fix and upload again: {message}" if message else "Fix and upload again"
        for message in messages
    ]
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        report.to_excel(writer, sheet_name=SHEET_NAME, index=False)
    return output.getvalue()


def import_asset_excel(
    file: BinaryIO | bytes,
    repository: AssetRepository | None = None,
    account_id: int | None = None,
) -> dict[str, Any]:
    validation = validate_asset_excel(file)
    if not validation.success:
        return {
            "success": False,
            "imported_count": 0,
            "backup_path": None,
            "errors": validation.errors,
            "warnings": validation.warnings,
        }

    target = repository or get_asset_repository()
    backup_path = target.backup()
    if account_id is None:
        target.replace_all_assets(validation.rows)
    else:
        target.replace_all_assets(validation.rows, account_id=account_id)

    return {
        "success": True,
        "imported_count": len(validation.rows),
        "backup_path": str(backup_path) if backup_path else None,
        "errors": [],
        "warnings": validation.warnings,
    }


def import_asset_rows(
    rows: list[dict[str, Any]],
    warnings: list[str] | None = None,
    repository: AssetRepository | None = None,
    account_id: int | None = None,
) -> dict[str, Any]:
    """Apply already validated and resolved Smart Import rows atomically."""
    if not rows:
        return {
            "success": False,
            "imported_count": 0,
            "backup_path": None,
            "errors": ["반영할 자산이 없습니다."],
            "warnings": warnings or [],
        }
    target = repository or get_asset_repository()
    backup_path = target.backup()
    if account_id is None:
        target.replace_all_assets(rows)
    else:
        target.replace_all_assets(rows, account_id=account_id)
    return {
        "success": True,
        "imported_count": len(rows),
        "backup_path": str(backup_path) if backup_path else None,
        "errors": [],
        "warnings": warnings or [],
    }
